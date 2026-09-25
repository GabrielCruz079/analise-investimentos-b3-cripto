# ============================================
# PROJETO: Análise de Investimentos B3 + Cripto
# Script: EDA + Gráficos Automáticos + Excel
# Autor: Gabriel Ramos Cruz
# Período: 2023–2024
# ============================================

import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import matplotlib.dates as mdates
import seaborn as sns
import pymysql
import numpy as np
import openpyxl 
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from matplotlib.patches import Patch
import os

# ============================================
# CONFIGURAÇÕES VISUAIS
# ============================================
plt.rcParams.update({
    'figure.facecolor': '#1e1e2e',
    'axes.facecolor':   '#1e1e2e',
    'axes.edgecolor':   '#444466',
    'axes.labelcolor':  '#ccccdd',
    'xtick.color':      '#ccccdd',
    'ytick.color':      '#ccccdd',
    'text.color':       '#ccccdd',
    'grid.color':       '#333355',
    'grid.linestyle':   '--',
    'grid.alpha':       0.4,
    'font.family':      'DejaVu Sans',
    'font.size':        11,
})

VERDE    = '#00d68f'
AZUL     = '#3b8bff'
ROXO     = '#a78bfa'
LARANJA  = '#f59e0b'
VERMELHO = '#ef4444'
AMARELO  = '#facc15'
CIANO    = '#06b6d4'
LIMA     = '#84cc16'
CORES    = [VERDE, AZUL, ROXO, LARANJA, VERMELHO, CIANO, '#ec4899', LIMA]

CORES_ATIVOS = {
    'PETR4':   VERDE,
    'VALE3':   AZUL,
    'ITUB4':   ROXO,
    'BBDC4':   LARANJA,
    'MGLU3':   VERMELHO,
    'AMER3':   '#ec4899',
    'BTC-USD': AMARELO,
    'ETH-USD': CIANO,
    'SOL-USD': LIMA,
}

os.makedirs('graficos', exist_ok=True)
os.makedirs('dados',    exist_ok=True)
print("📁 Pastas criadas: graficos/ e dados/")

# ============================================
# CONEXÃO
# ============================================
conn = pymysql.connect(
    host='localhost',
    user='root',
    password='',        # sua senha aqui
    database='investimentos',
    charset='utf8mb4'
)
print("✅ Conectado ao banco investimentos!\n")

# ============================================
# EXTRAÇÃO DOS DADOS
# ============================================

# Preços históricos completos
df = pd.read_sql("""
    SELECT
        p.data,
        a.ticker,
        a.nome,
        a.tipo,
        a.setor,
        a.moeda,
        p.preco_fechamento,
        p.preco_abertura,
        p.preco_maximo,
        p.preco_minimo,
        p.volume,
        p.variacao_pct
    FROM precos_historicos p
    JOIN ativos a ON p.ativo_id = a.ativo_id
    ORDER BY p.data, a.ticker
""", conn, parse_dates=['data'])

# Indicadores técnicos
df_ind = pd.read_sql("""
    SELECT
        i.data,
        a.ticker,
        a.tipo,
        i.retorno_diario,
        i.retorno_acumulado,
        i.volatilidade_30d,
        i.media_movel_7d,
        i.media_movel_30d,
        i.media_movel_90d
    FROM indicadores i
    JOIN ativos a ON i.ativo_id = a.ativo_id
    ORDER BY i.data, a.ticker
""", conn, parse_dates=['data'])

# Portfólio
df_port = pd.read_sql("""
    SELECT
        a.ticker,
        a.nome,
        a.tipo,
        p.peso_pct,
        p.valor_investido,
        p.data_entrada
    FROM portfolio p
    JOIN ativos a ON p.ativo_id = a.ativo_id
""", conn)

conn.close()
print(f"✅ Dados carregados:")
print(f"   Preços históricos : {len(df):,} registros")
print(f"   Indicadores       : {len(df_ind):,} registros")
print(f"   Ativos            : {df['ticker'].nunique()}\n")

# ============================================
# GRÁFICO 1 — RETORNO ACUMULADO POR TIPO
# ============================================
print("📊 Gerando gráfico 1 — Retorno Acumulado...")

fig, axes = plt.subplots(2, 1, figsize=(14, 10))
fig.suptitle('Retorno Acumulado 2023–2024', fontsize=15, fontweight='bold', color='white', y=0.98)

for ticker in df_ind[df_ind['tipo'] == 'Ação']['ticker'].unique():
    sub = df_ind[df_ind['ticker'] == ticker].dropna(subset=['retorno_acumulado'])
    if not sub.empty:
        axes[0].plot(sub['data'], sub['retorno_acumulado'],
                     label=ticker, color=CORES_ATIVOS.get(ticker, AZUL),
                     linewidth=2, alpha=0.9)
axes[0].axhline(0, color='white', linestyle='--', alpha=0.3)
axes[0].set_title('Ações B3', color='white', fontsize=12)
axes[0].set_ylabel('Retorno Acumulado (%)', color='#aaaacc')
axes[0].legend(facecolor='#2a2a3e', edgecolor='#444466', labelcolor='white', fontsize=9)
axes[0].grid(alpha=0.3)
axes[0].yaxis.set_major_formatter(mticker.FormatStrFormatter('%.0f%%'))
axes[0].xaxis.set_major_formatter(mdates.DateFormatter('%b/%y'))

for ticker in df_ind[df_ind['tipo'] == 'Cripto']['ticker'].unique():
    sub = df_ind[df_ind['ticker'] == ticker].dropna(subset=['retorno_acumulado'])
    if not sub.empty:
        axes[1].plot(sub['data'], sub['retorno_acumulado'],
                     label=ticker, color=CORES_ATIVOS.get(ticker, VERDE),
                     linewidth=2, alpha=0.9)
axes[1].axhline(0, color='white', linestyle='--', alpha=0.3)
axes[1].set_title('Criptomoedas', color='white', fontsize=12)
axes[1].set_ylabel('Retorno Acumulado (%)', color='#aaaacc')
axes[1].set_xlabel('Data', color='#aaaacc')
axes[1].legend(facecolor='#2a2a3e', edgecolor='#444466', labelcolor='white', fontsize=9)
axes[1].grid(alpha=0.3)
axes[1].yaxis.set_major_formatter(mticker.FormatStrFormatter('%.0f%%'))
axes[1].xaxis.set_major_formatter(mdates.DateFormatter('%b/%y'))

plt.tight_layout()
plt.savefig('graficos/01_retorno_acumulado.png', dpi=150, bbox_inches='tight', facecolor='#1e1e2e')
plt.close()
print("   ✅ graficos/01_retorno_acumulado.png")

# ============================================
# GRÁFICO 2 — RANKING DE RETORNO TOTAL
# ============================================
print("📊 Gerando gráfico 2 — Ranking de Retorno...")

retorno_final = (df_ind.groupby('ticker')
                       .apply(lambda x: x.sort_values('data').iloc[-1]['retorno_acumulado'])
                       .reset_index()
                       .rename(columns={0: 'retorno_final'})
                       .sort_values('retorno_final', ascending=True))

cores_ranking = [VERDE if v >= 0 else VERMELHO for v in retorno_final['retorno_final']]

fig, ax = plt.subplots(figsize=(11, 6))
bars = ax.barh(retorno_final['ticker'], retorno_final['retorno_final'],
               color=cores_ranking, edgecolor='none', height=0.6)
ax.axvline(0, color='white', linestyle='--', alpha=0.4)
for bar, val in zip(bars, retorno_final['retorno_final']):
    x_pos = val + 1 if val >= 0 else val - 1
    ha = 'left' if val >= 0 else 'right'
    ax.text(x_pos, bar.get_y() + bar.get_height() / 2,
            f'{val:.1f}%', va='center', fontsize=10, color='white', ha=ha)

ax.set_title('Ranking de Retorno Acumulado 2023–2024',
             fontsize=14, fontweight='bold', color='white', pad=15)
ax.set_xlabel('Retorno Acumulado (%)', color='#aaaacc')
ax.grid(axis='x', alpha=0.3)

legend = [Patch(facecolor=VERDE, label='Retorno positivo'),
          Patch(facecolor=VERMELHO, label='Retorno negativo')]
ax.legend(handles=legend, facecolor='#2a2a3e', edgecolor='#444466', labelcolor='white')
plt.tight_layout()
plt.savefig('graficos/02_ranking_retorno.png', dpi=150, bbox_inches='tight', facecolor='#1e1e2e')
plt.close()
print("   ✅ graficos/02_ranking_retorno.png")

# ============================================
# GRÁFICO 3 — VOLATILIDADE POR ATIVO
# ============================================
print("📊 Gerando gráfico 3 — Volatilidade...")

vol_media = (df_ind.groupby(['ticker'])['volatilidade_30d']
                   .mean()
                   .reset_index()
                   .sort_values('volatilidade_30d', ascending=False))

# Adicionar tipo
tipos_dict = df[['ticker', 'tipo']].drop_duplicates().set_index('ticker')['tipo'].to_dict()
vol_media['tipo'] = vol_media['ticker'].map(tipos_dict)

fig, ax = plt.subplots(figsize=(11, 5))
cores_vol = [AMARELO if tipos_dict.get(t) == 'Cripto' else AZUL for t in vol_media['ticker']]
bars = ax.bar(vol_media['ticker'], vol_media['volatilidade_30d'],
              color=cores_vol, edgecolor='none', width=0.6)
for bar, val in zip(bars, vol_media['volatilidade_30d']):
    ax.text(bar.get_x() + bar.get_width() / 2, val + 0.05,
            f'{val:.2f}%', ha='center', fontsize=9, color='white')

ax.set_title('Volatilidade Média (30 dias) por Ativo — 2023–2024',
             fontsize=13, fontweight='bold', color='white', pad=15)
ax.set_ylabel('Volatilidade (%)', color='#aaaacc')
ax.grid(axis='y', alpha=0.3)

legend = [Patch(facecolor=AZUL, label='Ações B3'),
          Patch(facecolor=AMARELO, label='Criptomoedas')]
ax.legend(handles=legend, facecolor='#2a2a3e', edgecolor='#444466', labelcolor='white')
plt.tight_layout()
plt.savefig('graficos/03_volatilidade.png', dpi=150, bbox_inches='tight', facecolor='#1e1e2e')
plt.close()
print("   ✅ graficos/03_volatilidade.png")

# ============================================
# GRÁFICO 4 — EVOLUÇÃO DO PREÇO NORMALIZADO
# ============================================
print("📊 Gerando gráfico 4 — Evolução de Preços Normalizada...")

fig, axes = plt.subplots(2, 1, figsize=(14, 10))
fig.suptitle('Evolução do Preço Normalizado (Base 100 = Jan/2023)',
             fontsize=14, fontweight='bold', color='white', y=0.98)

for tipo, ax in zip(['Ação', 'Cripto'], axes):
    titulo = 'Ações B3' if tipo == 'Ação' else 'Criptomoedas'
    for ticker in df[df['tipo'] == tipo]['ticker'].unique():
        sub = df[df['ticker'] == ticker].sort_values('data')
        if sub.empty:
            continue
        base = sub['preco_fechamento'].iloc[0]
        normalizado = (sub['preco_fechamento'] / base) * 100
        ax.plot(sub['data'], normalizado,
                label=ticker, color=CORES_ATIVOS.get(ticker, AZUL),
                linewidth=2, alpha=0.9)
    ax.axhline(100, color='white', linestyle='--', alpha=0.3, label='Base 100')
    ax.set_title(titulo, color='white', fontsize=12)
    ax.set_ylabel('Preço Normalizado', color='#aaaacc')
    ax.legend(facecolor='#2a2a3e', edgecolor='#444466', labelcolor='white', fontsize=9)
    ax.grid(alpha=0.3)
    ax.xaxis.set_major_formatter(mdates.DateFormatter('%b/%y'))

axes[1].set_xlabel('Data', color='#aaaacc')
plt.tight_layout()
plt.savefig('graficos/04_preco_normalizado.png', dpi=150, bbox_inches='tight', facecolor='#1e1e2e')
plt.close()
print("   ✅ graficos/04_preco_normalizado.png")

# ============================================
# GRÁFICO 5 — CORRELAÇÃO ENTRE ATIVOS
# ============================================
print("📊 Gerando gráfico 5 — Correlação...")

pivot_ret = df_ind.pivot_table(index='data', columns='ticker', values='retorno_diario')
corr = pivot_ret.corr()

fig, ax = plt.subplots(figsize=(11, 8))
sns.heatmap(corr, annot=True, fmt='.2f', cmap='RdYlGn',
            center=0, vmin=-1, vmax=1,
            linewidths=0.5, linecolor='#1e1e2e',
            annot_kws={'size': 9, 'color': 'black'},
            ax=ax)
ax.set_title('Correlação entre Ativos — Retornos Diários 2023–2024',
             fontsize=13, fontweight='bold', color='white', pad=15)
ax.tick_params(colors='white')
plt.tight_layout()
plt.savefig('graficos/05_correlacao.png', dpi=150, bbox_inches='tight', facecolor='#1e1e2e')
plt.close()
print("   ✅ graficos/05_correlacao.png")

# ============================================
# GRÁFICO 6 — RISCO X RETORNO
# ============================================
print("📊 Gerando gráfico 6 — Risco x Retorno...")

risco_retorno = pd.merge(retorno_final,
                         vol_media[['ticker', 'volatilidade_30d', 'tipo']],
                         on='ticker')

fig, ax = plt.subplots(figsize=(11, 7))
for _, row in risco_retorno.iterrows():
    cor = CORES_ATIVOS.get(row['ticker'], AZUL)
    ax.scatter(row['volatilidade_30d'], row['retorno_final'],
               color=cor, s=200, zorder=5, edgecolors='white', linewidths=0.8)
    ax.annotate(row['ticker'],
                (row['volatilidade_30d'], row['retorno_final']),
                textcoords='offset points', xytext=(8, 4),
                fontsize=10, color=cor, fontweight='bold')

ax.axhline(0, color='white', linestyle='--', alpha=0.3)
ax.axvline(risco_retorno['volatilidade_30d'].mean(), color=LARANJA,
           linestyle='--', alpha=0.4, label='Volatilidade média')
ax.set_title('Risco × Retorno — 2023–2024',
             fontsize=14, fontweight='bold', color='white', pad=15)
ax.set_xlabel('Volatilidade Média 30d (%)', color='#aaaacc')
ax.set_ylabel('Retorno Acumulado (%)', color='#aaaacc')
ax.legend(facecolor='#2a2a3e', edgecolor='#444466', labelcolor='white')
ax.grid(alpha=0.3)

ax.text(0.02, 0.97, 'Baixo Risco\nAlto Retorno', transform=ax.transAxes,
        fontsize=8, color=VERDE, va='top', alpha=0.6)
ax.text(0.75, 0.97, 'Alto Risco\nAlto Retorno', transform=ax.transAxes,
        fontsize=8, color=LARANJA, va='top', alpha=0.6)
ax.text(0.02, 0.05, 'Baixo Risco\nBaixo Retorno', transform=ax.transAxes,
        fontsize=8, color=AZUL, va='bottom', alpha=0.6)
ax.text(0.75, 0.05, 'Alto Risco\nBaixo Retorno', transform=ax.transAxes,
        fontsize=8, color=VERMELHO, va='bottom', alpha=0.6)

plt.tight_layout()
plt.savefig('graficos/06_risco_retorno.png', dpi=150, bbox_inches='tight', facecolor='#1e1e2e')
plt.close()
print("   ✅ graficos/06_risco_retorno.png")

# ============================================
# GRÁFICO 7 — PORTFÓLIO SIMULADO
# ============================================
print("📊 Gerando gráfico 7 — Portfólio Simulado...")

# Mapeia o retorno final de cada ativo (calculado no Gráfico 2) para o portfólio.
# fillna(0) evita erro caso algum ticker do portfólio não tenha indicador calculado.
retorno_dict = dict(zip(retorno_final['ticker'], retorno_final['retorno_final']))
df_port['retorno_pct']    = df_port['ticker'].map(retorno_dict).fillna(0)
df_port['valor_atual']    = df_port['valor_investido'] * (1 + df_port['retorno_pct'] / 100)
df_port['lucro_prejuizo'] = df_port['valor_atual'] - df_port['valor_investido']

# Filtrar apenas ativos com dados válidos para o pie
df_port_pie = df_port[df_port['valor_atual'].notna() & (df_port['valor_atual'] > 0)].copy()
cores_port = [CORES_ATIVOS.get(t, AZUL) for t in df_port_pie['ticker']]

fig, axes = plt.subplots(1, 2, figsize=(14, 6))
fig.suptitle('Portfólio Simulado — R$ 10.000 Investidos em Jan/2023',
             fontsize=13, fontweight='bold', color='white')

axes[0].pie(df_port_pie['valor_atual'], labels=df_port_pie['ticker'],
            colors=cores_port, autopct='%1.1f%%', startangle=90,
            wedgeprops={'edgecolor': '#1e1e2e', 'linewidth': 1.5},
            textprops={'color': 'white', 'fontsize': 9})
axes[0].set_title('Composição Atual do Portfólio', color='white')

cores_lp = [VERDE if v >= 0 else VERMELHO for v in df_port['lucro_prejuizo']]
bars = axes[1].bar(df_port['ticker'], df_port['lucro_prejuizo'],
                   color=cores_lp, edgecolor='none', width=0.6)
axes[1].axhline(0, color='white', linestyle='--', alpha=0.4)
for bar, val in zip(bars, df_port['lucro_prejuizo']):
    axes[1].text(bar.get_x() + bar.get_width() / 2,
                 val + (5 if val >= 0 else -15),
                 f'R$ {val:.0f}', ha='center', fontsize=8,
                 color='white', fontweight='bold')
axes[1].set_title('Lucro / Prejuízo por Ativo (R$)', color='white')
axes[1].set_ylabel('R$', color='#aaaacc')
axes[1].grid(axis='y', alpha=0.3)
plt.setp(axes[1].xaxis.get_majorticklabels(), rotation=30)

plt.tight_layout()
plt.savefig('graficos/07_portfolio.png', dpi=150, bbox_inches='tight', facecolor='#1e1e2e')
plt.close()
print("   ✅ graficos/07_portfolio.png")

# ============================================
# EXPORTAR EXCEL PARA POWER BI
# ============================================
print("\n📗 Exportando Excel para Power BI...")

ACCENT = '1a56a0'
BRANCO = 'FFFFFF'
CINZA  = 'f0f4ff'

def formatar_header(ws, cor_bg=ACCENT):
    for cell in ws[1]:
        cell.fill      = PatternFill('solid', fgColor=cor_bg)
        cell.font      = Font(bold=True, color=BRANCO, size=10, name='Calibri')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 28

def formatar_linhas(ws, cor_par=CINZA):
    lado = Side(style='thin', color='CCCCCC')
    borda = Border(left=lado, right=lado, top=lado, bottom=lado)
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.border    = borda
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font      = Font(size=9, name='Calibri')
            if row_idx % 2 == 0:
                cell.fill = PatternFill('solid', fgColor=cor_par)

def ajustar_cols(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 35)

# Preparar DataFrames
df_export = df.copy()
df_export['data'] = df_export['data'].dt.strftime('%Y-%m-%d')

df_ind_export = df_ind.copy()
df_ind_export['data'] = df_ind_export['data'].dt.strftime('%Y-%m-%d')

# Resumo por ativo
resumo = pd.merge(retorno_final,
                  vol_media[['ticker', 'volatilidade_30d', 'tipo']],
                  on='ticker')
resumo.columns = ['Ticker', 'Retorno Acumulado (%)', 'Volatilidade Média (%)', 'Tipo']
resumo['Sharpe Simples'] = (
    resumo['Retorno Acumulado (%)'] / resumo['Volatilidade Média (%)']
).round(2)
resumo = resumo.sort_values('Retorno Acumulado (%)', ascending=False)

# Portfólio exportação
df_port_exp = df_port[['ticker', 'tipo', 'peso_pct', 'valor_investido',
                         'retorno_pct', 'valor_atual', 'lucro_prejuizo']].copy()
df_port_exp.columns = ['Ticker', 'Tipo', 'Peso (%)', 'Valor Investido (R$)',
                        'Retorno (%)', 'Valor Atual (R$)', 'Lucro/Prejuízo (R$)']

arquivo = 'dados/investimentos_powerbi.xlsx'
with pd.ExcelWriter(arquivo, engine='openpyxl') as writer:
    df_export.to_excel(writer,      sheet_name='Preços Históricos', index=False)
    df_ind_export.to_excel(writer,  sheet_name='Indicadores',       index=False)
    resumo.to_excel(writer,         sheet_name='Resumo Ativos',     index=False)
    df_port_exp.to_excel(writer,    sheet_name='Portfólio',         index=False)

wb = openpyxl.load_workbook(arquivo)
cores_abas = {
    'Preços Históricos': '1a56a0',
    'Indicadores':       '16a085',
    'Resumo Ativos':     '8e44ad',
    'Portfólio':         'c0392b',
}
for aba, cor in cores_abas.items():
    ws = wb[aba]
    formatar_header(ws, cor)
    formatar_linhas(ws)
    ajustar_cols(ws)
    ws.freeze_panes = 'A2'

wb.save(arquivo)
print(f"   ✅ {arquivo}")

# ============================================
# RESUMO FINAL
# ============================================
total_investido = df_port['valor_investido'].sum()
total_atual     = df_port['valor_atual'].sum()
lucro_total     = total_atual - total_investido
retorno_port    = (lucro_total / total_investido) * 100
melhor          = resumo.iloc[0]
pior            = resumo.iloc[-1]

print(f"\n{'='*55}")
print(f"✅ ANÁLISE CONCLUÍDA — 7 GRÁFICOS + EXCEL GERADOS")
print(f"{'='*55}")
print(f"\n📊 RESUMO EXECUTIVO — Portfólio B3 + Cripto:")
print(f"   Valor investido          : R$ {total_investido:>12,.2f}")
print(f"   Valor atual              : R$ {total_atual:>12,.2f}")
print(f"   Lucro / Prejuízo         : R$ {lucro_total:>12,.2f}")
print(f"   Retorno total portfólio  :    {retorno_port:>11.2f}%")
print(f"   Melhor ativo             : {melhor['Ticker']} ({melhor['Retorno Acumulado (%)']:.1f}%)")
print(f"   Pior ativo               : {pior['Ticker']} ({pior['Retorno Acumulado (%)']:.1f}%)")
print(f"\n📁 Arquivos gerados:")
print(f"   graficos/ → 7 gráficos PNG")
print(f"   dados/investimentos_powerbi.xlsx → 4 abas para Power BI")
print(f"{'='*55}")